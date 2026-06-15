import json
from pathlib import Path
from src.content_engine import ContentEngine
from src.config import COURSES, BRAND


class SlideDeckGenerator:
    """Generates slide deck outlines and content for PPTX production."""

    SLIDE_TEMPLATES = {
        "title": ["title", "subtitle"],
        "content": ["title", "bullet_points"],
        "two_column": ["title", "left_points", "right_points"],
        "quote": ["quote", "attribution"],
        "cta": ["headline", "action", "url"],
    }

    def __init__(self, engine: ContentEngine):
        self.engine = engine
        self.output_dir = Path("outputs/slide_decks")
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(self, course_key: str) -> dict:
        course = COURSES[course_key]

        prompt = f"""Create a complete webinar slide deck outline for: {course['title']}

Course: {course['description']}

Design a 15–20 slide deck that:
1. Opens with a hook that stops scrollers
2. Establishes credibility for Ashanté Richardson / 1787 Wealth Academy
3. Teaches core concepts from the course
4. Builds desire for the full paid course
5. Closes with a clear CTA to join 1787 Wealth Academy on Skool

For each slide provide:
- slide_number: int
- type: "title" | "content" | "two_column" | "quote" | "cta" | "section_break"
- title: slide title
- content: list of bullet points OR a paragraph (depending on type)
- speaker_notes: what Ashanté says on this slide (2-4 sentences)
- design_notes: color/visual suggestions using navy #1A3A5C and gold #C9A84C

Return as JSON with:
{{
  "deck_title": "...",
  "course": "...",
  "use_case": "webinar/workshop",
  "slides": [...]
}}"""

        response = self.engine.generate_long(prompt)
        start = response.find("{")
        end = response.rfind("}") + 1
        deck = json.loads(response[start:end])

        output_path = self.output_dir / f"{course_key}_slide_deck.json"
        with open(output_path, "w") as f:
            json.dump(deck, f, indent=2)

        md_path = self.output_dir / f"{course_key}_slide_deck.md"
        self._write_markdown(deck, md_path)

        print(f"  ✓ {len(deck.get('slides', []))} slides → {output_path}")
        return deck

    def _write_markdown(self, deck: dict, path: Path):
        lines = [
            f"# Slide Deck: {deck.get('deck_title', '')}",
            f"*Course: {deck.get('course', '')}*",
            f"*Use case: {deck.get('use_case', '')}*\n",
            "---\n",
        ]
        for slide in deck.get("slides", []):
            n = slide.get("slide_number", "?")
            stype = slide.get("type", "")
            title = slide.get("title", "")
            lines.append(f"## Slide {n} [{stype.upper()}]: {title}")

            content = slide.get("content", [])
            if isinstance(content, list):
                for point in content:
                    lines.append(f"- {point}")
            elif isinstance(content, str):
                lines.append(content)

            if slide.get("speaker_notes"):
                lines.append(f"\n**Speaker notes:** {slide['speaker_notes']}")
            if slide.get("design_notes"):
                lines.append(f"**Design:** {slide['design_notes']}")
            lines.append("\n---\n")

        path.write_text("\n".join(lines))

    def generate_pptx(self, course_key: str, deck: dict):
        """Build a branded .pptx file from a generated deck outline."""
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
        except ImportError:
            print("  ⚠ python-pptx not installed — skipping PPTX generation (pip install python-pptx)")
            return

        prs = Presentation()
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        navy = RGBColor(0x1A, 0x3A, 0x5C)
        gold = RGBColor(0xC9, 0xA8, 0x4C)

        blank_layout = prs.slide_layouts[6]

        for slide_data in deck.get("slides", []):
            slide = prs.slides.add_slide(blank_layout)

            # Background
            background = slide.background
            fill = background.fill
            fill.solid()
            fill.fore_color.rgb = navy

            title_text = slide_data.get("title", "")
            if title_text:
                txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1.2))
                tf = txBox.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = title_text
                p.font.size = Pt(36)
                p.font.bold = True
                p.font.color.rgb = gold

            content = slide_data.get("content", [])
            if isinstance(content, list) and content:
                txBox = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(12), Inches(4.5))
                tf = txBox.text_frame
                tf.word_wrap = True
                for i, point in enumerate(content):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = f"• {point}"
                    p.font.size = Pt(22)
                    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif isinstance(content, str) and content:
                txBox = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(12), Inches(4.5))
                tf = txBox.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = content
                p.font.size = Pt(22)
                p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        pptx_path = self.output_dir / f"{course_key}_slide_deck.pptx"
        prs.save(str(pptx_path))
        print(f"  ✓ PPTX saved → {pptx_path}")
